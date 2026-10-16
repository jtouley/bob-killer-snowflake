"""Shared builders for the extract tests.

Structure fixtures (sheets, formula text, names, tables, detectors) are authored here with
openpyxl at test time, so no binary fixture is committed. openpyxl cannot calculate, so these
workbooks carry no cached formula values. Cached values come only from the Excel-saved golden
workbooks, whose expected values live in ``tests/oracle/extract_golden.json``.

Shared Workbook IR builders for the IR and generator tests.

The IR is written as its JSON/YAML shape and validated through the strict contract, so the
builders document the wire shape too. One IR holds every node kind (core SPEC "Workbook IR").
"""

from __future__ import annotations

import hashlib
import io
import json
import zipfile
from collections.abc import Callable
from pathlib import Path
from typing import TYPE_CHECKING, Any

import openpyxl
import pytest
from openpyxl.chart import BarChart, Reference
from openpyxl.packaging.relationship import Relationship
from openpyxl.pivot.cache import CacheDefinition, CacheField, CacheSource, WorksheetSource
from openpyxl.pivot.record import RecordList
from openpyxl.pivot.table import Location, TableDefinition
from openpyxl.styles import PatternFill
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.workbook.external_link.external import ExternalBook, ExternalLink
from openpyxl.worksheet.formula import DataTableFormula
from openpyxl.worksheet.table import Table

if TYPE_CHECKING:
    from bob_killer.contracts.raw import RawWorkbook
from collections.abc import Mapping

if TYPE_CHECKING:
    from bob_killer.contracts.ir import CellRef, CellValue, TargetConfig, WorkbookIR


ORACLE_PATH = Path(__file__).resolve().parents[1] / "oracle" / "extract_golden.json"
ORACLE: dict[str, Any] = json.loads(ORACLE_PATH.read_text(encoding="utf-8"))


class Structure:
    """What the structure workbook holds; tests compare extraction against these inputs."""

    sheets = (("Inputs", "visible"), ("Calc", "visible"), ("Lists", "hidden"))
    # cell -> (value written, expected cached type, expected canonical cached text)
    constants = {
        ("Inputs", "A1"): ("qty", "string", "qty"),
        ("Inputs", "A2"): (3, "number", "3.0"),
        ("Inputs", "A3"): (0.1, "number", "0.1"),
        ("Inputs", "B1"): (True, "bool", "TRUE"),
        ("Lists", "A1"): ("north", "string", "north"),
    }
    formulas = {
        ("Calc", "A1"): "=Inputs!A2*2",
        ("Calc", "A2"): "=SUM(Inputs!A2:A3)",
        ("Calc", "A3"): '=IF(Inputs!B1,"yes","no")',
    }
    styled_blank = ("Inputs", "D5")
    plain_constant = ("Inputs", "A2")
    names = (
        ("", "rate", "Inputs!$A$3"),
        ("Calc", "local_total", "Calc!$A$2"),
    )
    tables = (("Prices", "Inputs", "F1:G3"),)


def structure_workbook() -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    for i, (name, state) in enumerate(Structure.sheets):
        ws = wb.active if i == 0 else wb.create_sheet()
        assert ws is not None
        ws.title = name
        ws.sheet_state = state  # type: ignore[assignment]
    for (sheet, address), (value, _, _) in Structure.constants.items():
        wb[sheet][address] = value
    for (sheet, address), formula in Structure.formulas.items():
        wb[sheet][address] = formula
    sheet, address = Structure.styled_blank
    wb[sheet][address].fill = PatternFill("solid", fgColor="FFFF00")
    for scope, name, refers_to in Structure.names:
        target = wb.defined_names if not scope else wb[scope].defined_names
        target[name] = DefinedName(name, attr_text=refers_to)
    for name, sheet, ref in Structure.tables:
        ws = wb[sheet]
        ws["F1"], ws["G1"] = "region", "price"
        ws.add_table(Table(displayName=name, ref=ref))
    return wb


def build_structure() -> bytes:
    return save(structure_workbook())


def save(wb: openpyxl.Workbook, extra_parts: dict[str, bytes] | None = None) -> bytes:
    buf = io.BytesIO()
    wb.save(buf)
    if extra_parts:
        with zipfile.ZipFile(buf, "a") as zf:
            for name, data in extra_parts.items():
                zf.writestr(name, data)
    return buf.getvalue()


# Unsupported-feature kinds: the names detectors register under, plus the VBA inventory's.
CHART = "chart"
PIVOT_TABLE = "pivot_table"
EXTERNAL_LINK = "external_link"
DATA_TABLE = "data_table"
VBA_SOURCE = "vba_source"
VBA_PROJECT = "vba_project"
VBA_PART = "xl/vbaProject.bin"


def _add_chart(wb: openpyxl.Workbook) -> str:
    ws = wb["Inputs"]
    chart = BarChart()
    chart.add_data(Reference(ws, min_col=1, min_row=1, max_row=3))
    ws.add_chart(chart, "J5")
    return "Inputs!J5"


def _add_pivot_table(wb: openpyxl.Workbook) -> str:
    source = WorksheetSource(ref="A1:A3", sheet="Inputs")
    cache = CacheDefinition(
        cacheSource=CacheSource(type="worksheet", worksheetSource=source),
        cacheFields=[CacheField(name="qty")],
    )
    cache.records = RecordList()
    location = Location(ref="L1:M3", firstHeaderRow=1, firstDataRow=1, firstDataCol=1)
    pivot = TableDefinition(name="P1", cacheId=1, dataCaption="Values", location=location)
    pivot.cache = cache
    wb["Calc"]._pivots.append(pivot)
    return "Calc!L1:M3"


def _add_external_link(wb: openpyxl.Workbook) -> str:
    link = ExternalLink(externalBook=ExternalBook(id="rId1"))
    link.file_link = Relationship(
        type="externalLinkPath", Target="other.xlsx", TargetMode="External"
    )
    wb._external_links.append(link)
    return "[1]"


def _add_data_table(wb: openpyxl.Workbook) -> str:
    wb["Calc"]["C1"] = DataTableFormula(ref="C1:C3", dt2D=False, dtr=False, r1="A1")
    return "Calc!C1:C3"


FEATURES: dict[str, Callable[[openpyxl.Workbook], str]] = {
    CHART: _add_chart,
    PIVOT_TABLE: _add_pivot_table,
    EXTERNAL_LINK: _add_external_link,
    DATA_TABLE: _add_data_table,
}


def build_with_feature(kind: str) -> tuple[bytes, str]:
    """The structure workbook plus one unsupported feature; returns (bytes, expected origin)."""
    wb = structure_workbook()
    origin = FEATURES[kind](wb)
    return save(wb), origin


def build_with_vba_part(data: bytes) -> bytes:
    return save(structure_workbook(), {VBA_PART: data})


def golden_path(workbook_id: str) -> Path:
    """The golden file pinned by the oracle: golden/ first, else fetched and hash-checked."""
    import fetch_golden
    from bob_killer.contracts.golden import GoldenStatus

    sha256 = ORACLE["workbooks"][workbook_id]["sha256"]
    entry = next(
        e for e in fetch_golden.load_lock(fetch_golden.LOCK_PATH).workbooks if e.id == workbook_id
    )
    assert entry.filename is not None
    path = fetch_golden.DEST / entry.filename
    if not path.is_file() or hashlib.sha256(path.read_bytes()).hexdigest() != sha256:
        pinned = entry.model_copy(update={"sha256": sha256, "status": GoldenStatus.RESOLVED})
        fetch_golden.fetch_entry(pinned, fetch_golden.DEST, accept=False)
    return path


@pytest.fixture(scope="session")
def golden_raw() -> Callable[[str], RawWorkbook]:
    from bob_killer.extract import extract_workbook

    cache: dict[str, RawWorkbook] = {}

    def get(workbook_id: str) -> RawWorkbook:
        if workbook_id not in cache:
            cache[workbook_id] = extract_workbook(golden_path(workbook_id).read_bytes())
        return cache[workbook_id]

    return get


@pytest.fixture(scope="session")
def structure_raw() -> RawWorkbook:
    from bob_killer.extract import extract_workbook

    return extract_workbook(build_structure())


Json = dict[str, Any]
IrMutation = Callable[[Json], None]


def col_letters(col: int) -> str:
    """Excel column letters for a 1-based column number (independent of the contract parser)."""
    letters = ""
    while col:
        col, rem = divmod(col - 1, 26)
        letters = chr(ord("A") + rem) + letters
    return letters


def a1(
    sheet: str, row: int, col: int, last_row: int | None = None, last_col: int | None = None
) -> str:
    quoted = "'" + sheet.replace("'", "''") + "'"
    head = f"{quoted}!{col_letters(col)}{row}"
    if last_row is None or last_col is None:
        return head
    return f"{head}:{col_letters(last_col)}{last_row}"


def cell(sheet: str, row: int, col: int, *, relative: bool = False) -> Json:
    return {
        "type": "cell",
        "sheet": sheet,
        "row": row,
        "col": col,
        "row_abs": not relative,
        "col_abs": not relative,
    }


def rng(sheet: str, row: int, col: int, last_row: int, last_col: int) -> Json:
    return {
        "type": "range",
        "sheet": sheet,
        "row": row,
        "col": col,
        "last_row": last_row,
        "last_col": last_col,
    }


def num(x: float) -> Json:
    return {"k": "n", "h": float(x).hex()}


def text(s: str) -> Json:
    return {"k": "s", "t": s}


def lit(value: Json) -> Json:
    return {"type": "lit", "value": value}


def param(pid: str) -> Json:
    return {"type": "param", "id": pid}


def call(name: str, *args: Json) -> Json:
    return {"type": "call", "name": name, "args": list(args)}


def op(name: str, *args: Json) -> Json:
    return {"type": "op", "op": name, "args": list(args)}


def node(nid: str, label: str, origin: str, *, status: str = "accepted") -> Json:
    return {
        "id": nid,
        "label": label,
        "origin": [origin],
        "description": f"{label} from {origin}",
        "name_rule": "defined_name",
        "name_status": status,
    }


def column(name: str, dtype: str) -> Json:
    return {"name": name, "label": name.replace("_", " ").title(), "dtype": dtype}


def expr_step(kind: str, nid: str, origin: str, expr: Json, formula: str, **extra: Any) -> Json:
    return {
        **node(nid, nid.replace("_", " ").title(), origin),
        "kind": kind,
        "expr": expr,
        "formula": formula,
        "fingerprint": formula,
        **extra,
    }


DATA, SUMMARY, RATES = "Data", "Summary", "Rates {2013}"


def sample_ir_json() -> Json:
    """Every IR section and every step kind, on three sheets (one name needs SQL/Jinja escaping)."""
    return {
        "schema_version": 1,
        "source": {"file": "bobs_forecast_2013.xlsm", "sha256": "9f2c" * 16},
        "sheets": [
            {"id": "data", "name": DATA, "index": 0},
            {"id": "summary", "name": SUMMARY, "index": 1},
            {"id": "rates", "name": RATES, "index": 2},
        ],
        "parameters": [
            {
                **node("growth_rate", "Annual growth rate", a1(SUMMARY, 4, 3)),
                "format": "percent",
                "default": num(0.035),
            }
        ],
        "inputs": [
            {
                **node("units_sold", "Units sold", a1(DATA, 2, 1, 5, 3)),
                "shape": "table",
                "kind": "input",
                "columns": [
                    column("region", "text"),
                    column("units", "float64"),
                    column("price", "float64"),
                ],
            },
            {
                **node("tax_rate", "Tax rate", a1(SUMMARY, 2, 2)),
                "shape": "scalar",
                "kind": "input",
                "columns": [column("value", "float64")],
            },
            {
                **node("today", "Today", a1(SUMMARY, 3, 2)),
                "shape": "scalar",
                "kind": "frozen_volatile",
                "columns": [column("value", "float64")],
            },
        ],
        "lookups": [
            {
                **node("region_tax", "Region tax", a1(RATES, 2, 1, 4, 2)),
                "key": "region",
                "columns": [column("region", "text"), column("rate", "float64")],
            }
        ],
        "steps": [
            expr_step(
                "column_expr",
                "revenue",
                a1(DATA, 2, 4, 5, 4),
                op("MUL", cell(DATA, 2, 2, relative=True), cell(DATA, 2, 3, relative=True)),
                "=B2*C2",
                over="units_sold",
            ),
            expr_step(
                "aggregate",
                "total_units",
                a1(SUMMARY, 5, 2),
                call("SUM", rng(DATA, 2, 2, 5, 2)),
                "=SUM(Data!B2:B5)",
            ),
            expr_step(
                "aggregate",
                "total_revenue",
                a1(SUMMARY, 6, 2),
                call("SUM", rng(DATA, 2, 4, 5, 4)),
                "=SUM(Data!D2:D5)",
            ),
            expr_step(
                "scalar_expr",
                "grown",
                a1(SUMMARY, 7, 2),
                op("MUL", cell(SUMMARY, 6, 2), op("ADD", lit(num(1)), param("growth_rate"))),
                "=B6*(1+0.035)",
            ),
            expr_step(
                "scalar_expr",
                "tax",
                a1(SUMMARY, 8, 2),
                op("MUL", cell(SUMMARY, 7, 2), cell(SUMMARY, 2, 2)),
                "=B7*B2",
            ),
            expr_step(
                "lookup",
                "east_rate",
                a1(SUMMARY, 9, 2),
                call(
                    "VLOOKUP",
                    lit(text("east {{x}}'")),
                    rng(RATES, 2, 1, 4, 2),
                    lit(num(2)),
                    lit({"k": "b", "b": False}),
                ),
                "=VLOOKUP(\"east\",'Rates {2013}'!A2:B4,2,FALSE)",
                lookup="region_tax",
            ),
            expr_step(
                "conditional",
                "band",
                a1(SUMMARY, 10, 2),
                call(
                    "IF",
                    op("GT", cell(SUMMARY, 8, 2), lit(num(100))),
                    lit(text("high")),
                    lit(text("low")),
                ),
                '=IF(B8>100,"high","low")',
            ),
            {
                **node("refresh", "Refresh totals", a1(SUMMARY, 11, 2)),
                "kind": "vba_procedure",
                "module": "Module1",
                "procedure": "refresh_totals",
                "status": "translated_verified",
                "assignments": [
                    {
                        "cell": {"sheet": SUMMARY, "row": 11, "col": 2},
                        "expr": op("ADD", cell(SUMMARY, 7, 2), cell(SUMMARY, 3, 2)),
                    }
                ],
            },
        ],
        "outputs": [
            {
                **node("tax_owed", "Total tax owed", a1(SUMMARY, 8, 2)),
                "step": "tax",
                "format": "currency",
            },
            {
                **node("revenue_by_row", "Revenue", a1(DATA, 2, 4, 5, 4)),
                "step": "revenue",
                "format": "number",
            },
        ],
        "unsupported": [
            {
                "kind": "pivot_table",
                "origin": a1(SUMMARY, 1, 4),
                "reason": "pivot tables are out of scope for v1",
            }
        ],
        "vba": [
            {
                "module": "Module1",
                "procedure": "refresh_totals",
                "status": "translated_verified",
                "classification": "pure_computation",
                "writes": [a1(SUMMARY, 11, 2)],
                "translated_to": "refresh",
            },
            {
                "module": "Module1",
                "procedure": "go_home",
                "status": "untranslated",
                "classification": "navigation_ui",
                "writes": [],
                "translated_to": None,
            },
        ],
    }


def build_ir(data: Mapping[str, Any]) -> WorkbookIR:
    from bob_killer.contracts.ir import WorkbookIR

    return WorkbookIR.model_validate_json(json.dumps(data))


def sample_ir(*mutations: IrMutation) -> WorkbookIR:
    data = sample_ir_json()
    for mutate in mutations:
        mutate(data)
    return build_ir(data)


def sample_target(run_id: str = "0123456789abcdef0123456789abcdef") -> TargetConfig:
    from bob_killer.contracts.ir import TargetConfig

    return TargetConfig(
        run_id=run_id,
        database="BOB_KILLER",
        deploy_schema="TEST_ABCD1234",
        role="BOB_KILLER_OWNER",
        warehouse="BOB_KILLER_WH",
        dbt_version="1.12.3",
    )


def sample_cells(ir: WorkbookIR) -> dict[CellRef, CellValue]:
    """A deterministic value for every cell an input, lookup or step covers (index-derived)."""
    from bob_killer.contracts.ir import CellValue, cell_range

    origins = [o for group in (ir.inputs, ir.lookups, ir.steps) for n in group for o in n.origin]
    values: dict[CellRef, CellValue] = {}
    for i, ref in enumerate(r for o in origins for r in cell_range(o).cells()):
        values[ref] = CellValue.text(f"t{i} {{{{ '") if i % 3 == 0 else CellValue.number(i / 7 - 1)
    return values
