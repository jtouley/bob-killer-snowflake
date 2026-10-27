"""Write the bronze landing fixtures with openpyxl (docs/SPEC.md §8 Phase 5). This script is the source.

Bronze reads literal cell values only, so the fixtures hold no formulas and need no Excel cached
values. Expected rows live in ``tests/oracle/bronze_<name>.json``.

    uv run python tests/fixtures/make_bronze_fixtures.py   # rewrites tests/fixtures/workbooks/bronze_*.xlsx
"""

from __future__ import annotations

from collections.abc import Callable
from datetime import datetime, time, timedelta
from pathlib import Path

import openpyxl

WORKBOOKS = Path(__file__).resolve().parent / "workbooks"

DATA_SHEET = "Data"
SECOND_SHEET = "Second"
EMPTY_SHEET = "Empty"


def basic() -> openpyxl.Workbook:
    """Two sheets with data, one empty sheet.

    Data: rows 1-2 above the header are empty; row 3 is the header with a blank and a duplicate
    name; row 5 is a blank data row (None and empty text); row 6 has an error, a date, a time and
    a duration; row 7 has blanks between values.
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    assert ws is not None
    ws.title = DATA_SHEET
    ws["A3"], ws["C3"], ws["D3"], ws["E3"] = "id", "name", "id", "when"
    ws["A4"], ws["B4"], ws["C4"], ws["D4"] = 1, 2.5, "alpha", True
    ws["A5"], ws["C5"] = "", None
    ws["A6"], ws["B6"] = 2, "#DIV/0!"
    ws["C6"], ws["D6"] = datetime(2026, 1, 2, 3, 4, 5), time(13, 30)
    ws["E6"] = timedelta(hours=1, minutes=30)
    ws["A7"], ws["D7"] = 3, False

    second = wb.create_sheet(SECOND_SHEET)
    second.append(["x", "x", "x"])
    second.append([10, None, "z"])
    wb.create_sheet(EMPTY_SHEET)
    return wb


FIXTURES: dict[str, Callable[[], openpyxl.Workbook]] = {"bronze_basic": basic}


def write(name: str, directory: Path = WORKBOOKS) -> Path:
    directory.mkdir(parents=True, exist_ok=True)
    path = directory / f"{name}.xlsx"
    FIXTURES[name]().save(path)
    return path


def main() -> int:
    for name in FIXTURES:
        print(write(name))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
