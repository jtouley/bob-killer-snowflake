"""Read a workbook twice with openpyxl: formula text, then Excel's cached values.

Nothing is evaluated or executed: formulas stay text, and a value is only ever what Excel saved.
"""

from __future__ import annotations

import hashlib
import io
import zipfile
from collections.abc import Iterator
from datetime import date, datetime, time

import openpyxl
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.formula import ArrayFormula
from openpyxl.worksheet.worksheet import Worksheet

from bob_killer.contracts.raw import (
    CachedType,
    RawCell,
    RawDefinedName,
    RawSheet,
    RawSource,
    RawTable,
    RawWorkbook,
    SheetState,
)
from bob_killer.extract import vba
from bob_killer.extract.detectors import detect

# openpyxl keeps t="str" on a formula cell whose cached result is the empty string.
_EMPTY_STRING_RESULT = "str"


def load(data: bytes, *, data_only: bool) -> Workbook:
    return openpyxl.load_workbook(io.BytesIO(data), data_only=data_only)


def extract_workbook(data: bytes) -> RawWorkbook:
    formulas = load(data, data_only=False)
    values = load(data, data_only=True)
    with zipfile.ZipFile(io.BytesIO(data)) as zf:
        vba_modules, vba_unsupported = vba.inventory(zf)
    return RawWorkbook(
        source=RawSource(sha256=hashlib.sha256(data).hexdigest()),
        sheets=tuple(
            RawSheet(name=ws.title, position=i, state=SheetState(ws.sheet_state))
            for i, ws in enumerate(formulas.worksheets)
        ),
        cells=tuple(c for ws in formulas.worksheets for c in _cells(ws, values[ws.title])),
        defined_names=tuple(_defined_names(formulas)),
        tables=tuple(
            RawTable(name=t.displayName, sheet=ws.title, ref=t.ref)
            for ws in formulas.worksheets
            for t in ws.tables.values()
        ),
        vba_modules=vba_modules,
        unsupported=detect(formulas) + vba_unsupported,
    )


def _cells(formulas: Worksheet, values: Worksheet) -> Iterator[RawCell]:
    for key in sorted(formulas._cells):
        cell = formulas._cells[key]
        formula = formula_text(cell)
        cached_type, cached_value = _cached(values._cells[key], has_formula=bool(formula))
        yield RawCell(
            sheet=formulas.title,
            address=cell.coordinate,
            formula=formula,
            cached_value=cached_value,
            cached_type=cached_type,
            style_id=cell.style_id,
        )


def formula_text(cell: Cell | MergedCell) -> str:
    """The formula as written; "" for a constant. Data tables are reported, not kept as text."""
    if cell.data_type != "f":
        return ""
    value = cell.value
    if isinstance(value, ArrayFormula):
        return value.text or ""
    return value if isinstance(value, str) else ""


def _cached(cell: Cell | MergedCell, *, has_formula: bool) -> tuple[CachedType, str]:
    value = cell.value
    if value is None:
        if cell.data_type == _EMPTY_STRING_RESULT:
            return CachedType.STRING, ""
        return (CachedType.MISSING if has_formula else CachedType.BLANK), ""
    if isinstance(value, bool):
        return CachedType.BOOL, "TRUE" if value else "FALSE"
    if isinstance(value, int | float):
        return CachedType.NUMBER, repr(float(value))
    if isinstance(value, datetime | date | time):
        return CachedType.DATE, value.isoformat()
    if cell.data_type == "e":
        return CachedType.ERROR, str(value)
    return CachedType.STRING, str(value)


def _defined_names(wb: Workbook) -> Iterator[RawDefinedName]:
    for name, dn in wb.defined_names.items():
        yield RawDefinedName(scope="", name=name, refers_to=dn.attr_text or "")
    for ws in wb.worksheets:
        for name, dn in ws.defined_names.items():
            yield RawDefinedName(scope=ws.title, name=name, refers_to=dn.attr_text or "")
