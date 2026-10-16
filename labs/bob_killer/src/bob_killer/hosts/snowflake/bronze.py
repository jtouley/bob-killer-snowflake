"""Bronze landing: workbook rows as VARIANT payloads, once per file SHA-256 (SPEC §5.2, §6.1).

Bronze reads cached values only (openpyxl ``read_only=True, data_only=True``, D10). It is a
landing copy, never a parity source.
"""

from __future__ import annotations

import io
from collections.abc import Callable, Iterator, Sequence
from contextlib import contextmanager
from contextvars import ContextVar
from dataclasses import dataclass
from datetime import date, datetime, time, timedelta
from typing import Any

import openpyxl
from openpyxl.utils import get_column_letter

from bob_killer import registry
from bob_killer.contracts.bronze import BronzePayload, BronzeRow, CellJson, LandedFile
from bob_killer.hosts.snowflake.store import (
    ROWS_INSERTED,
    SnowflakeRunStore,
    SqlSession,
    affected,
    run_sql,
)
from bob_killer.store import RunRepository

LANDED = "landed"
ALREADY_LANDED = "already landed"
STAGE_NAME = "land_bronze"
UPLOADS = "@UPLOADS"
DEFAULT_BATCH_SIZE = 10_000
PARSE_JSON = "PARSE_JSON"

# openpyxl data_type values bronze keeps; an inline string is a string.
_DTYPES = {"inlineStr": "s", "str": "s", "f": "s"}


def _header(cells: Sequence[Any]) -> tuple[str, ...]:
    """A blank name is ``column_<letter>``; a repeat is ``<name>_2``, ``<name>_3``."""
    names: list[str] = []
    seen: dict[str, int] = {}
    for i, cell in enumerate(cells, start=1):
        base = f"column_{get_column_letter(i)}" if _blank(cell.value) else str(cell.value)
        seen[base] = seen.get(base, 0) + 1
        names.append(base if seen[base] == 1 else f"{base}_{seen[base]}")
    return tuple(names)


def _blank(value: object) -> bool:
    return value is None or value == ""


def _cell(cell: Any) -> tuple[CellJson, str]:
    value = cell.value
    if value is None:
        return None, "n"
    if isinstance(value, timedelta):
        return value.total_seconds(), "d"
    if isinstance(value, datetime | date | time):
        return value.isoformat(), "d"
    dtype: str = cell.data_type
    return value, _DTYPES.get(dtype, dtype)


def _sheet_rows(ws: Any, file_sha256: str, run_id: str) -> Iterator[BronzeRow]:
    header: tuple[str, ...] | None = None
    for row_number, cells in enumerate(ws.iter_rows(min_row=1), start=1):
        if all(_blank(c.value) for c in cells):
            continue
        if header is None:
            header = _header(cells)
            continue
        values = [_cell(c) for c in cells[: len(header)]]
        values += [(None, "n")] * (len(header) - len(values))
        yield BronzeRow(
            file_sha256=file_sha256,
            sheet=ws.title,
            row_number=row_number,
            run_id=run_id,
            payload=BronzePayload(
                column_names=header,
                dtypes={name: dtype for name, (_, dtype) in zip(header, values, strict=True)},
                data={name: value for name, (value, _) in zip(header, values, strict=True)},
            ),
        )


def rows_from_workbook(data: bytes, *, file_sha256: str, run_id: str) -> tuple[BronzeRow, ...]:
    """Each sheet's rows below its header; blank rows write nothing, row numbers do not shift."""
    wb = openpyxl.load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    try:
        return tuple(r for ws in wb.worksheets for r in _sheet_rows(ws, file_sha256, run_id))
    finally:
        wb.close()


def _landed_merge() -> str:
    names = tuple(LandedFile.model_fields)
    source = ", ".join(f"? AS {n}" for n in names)
    on = " AND ".join(f"t.{k} = s.{k}" for k in LandedFile.primary_key)
    return (
        f"MERGE INTO {LandedFile.table_name.upper()} t USING (SELECT {source}) s ON {on} "
        f"WHEN NOT MATCHED THEN INSERT ({', '.join(names)}) "
        f"VALUES ({', '.join(f's.{n}' for n in names)})"
    )


def _batch_insert(size: int) -> str:
    names = tuple(BronzeRow.model_fields)
    width = len(names)
    cols = [f"column{i}" for i in range(1, width + 1)]
    cols[-1] = f"{PARSE_JSON}({cols[-1]})"  # payload is the last field: JSON text to VARIANT
    marks = ", ".join(["(" + ", ".join("?" * width) + ")"] * size)
    return (
        f"INSERT INTO {BronzeRow.table_name.upper()} ({', '.join(names)}) "
        f"SELECT {', '.join(cols)} FROM VALUES {marks}"
    )


def _binds(row: BronzeRow) -> list[object]:
    return [row.file_sha256, row.sheet, row.row_number, row.run_id, row.payload.model_dump_json()]


class BronzeWriter:
    def __init__(self, session: SqlSession, *, batch_size: int = DEFAULT_BATCH_SIZE) -> None:
        self._session = session
        self.batch_size = batch_size

    def land(
        self,
        landed: LandedFile,
        rows: Sequence[BronzeRow],
        *,
        after_claim: Callable[[], None] | None = None,
    ) -> str:
        """Land the file once: BEGIN, MERGE INTO LANDED_FILES, batches, COMMIT; ROLLBACK on error.

        after_claim is a test seam for the concurrency live test; production callers pass nothing.
        """
        run_sql(self._session, "BEGIN")
        try:
            merged = run_sql(
                self._session, _landed_merge(), list(landed.model_dump(mode="json").values())
            )
            if affected(merged, ROWS_INSERTED) == 0:
                run_sql(self._session, "COMMIT")
                return ALREADY_LANDED
            if after_claim is not None:
                after_claim()
            for start in range(0, len(rows), self.batch_size):
                batch = rows[start : start + self.batch_size]
                binds = [v for r in batch for v in _binds(r)]
                run_sql(self._session, _batch_insert(len(batch)), binds)
            run_sql(self._session, "COMMIT")
        except BaseException:
            run_sql(self._session, "ROLLBACK")
            raise
        return LANDED


def land_file(
    session: SqlSession, landed: LandedFile, data: bytes, *, batch_size: int = DEFAULT_BATCH_SIZE
) -> str:
    """Parse the staged bytes and land their rows once (the BronzeWriter path of land_bronze)."""
    rows = rows_from_workbook(data, file_sha256=landed.file_sha256, run_id=landed.first_run_id)
    return BronzeWriter(session, batch_size=batch_size).land(landed, rows)


@dataclass(frozen=True)
class BronzeContext:
    """What land_bronze needs per procedure call (D4): the session and how to read a staged file.

    ``store`` and ``land`` default to SnowflakeRunStore and ``land_file`` over the session; the
    procedure host passes its own.
    """

    session: SqlSession
    read_file: Callable[[str], bytes]
    batch_size: int = DEFAULT_BATCH_SIZE
    store: RunRepository | None = None
    land: Callable[[LandedFile, bytes], str] | None = None


def stage_reader(session: Any) -> Callable[[str], bytes]:
    """Read ``@UPLOADS/<stage_file>`` with ``session.file.get_stream``."""

    def read(stage_file: str) -> bytes:
        data: bytes = session.file.get_stream(f"{UPLOADS}/{stage_file}").read()
        return data

    return read


_context: ContextVar[BronzeContext] = ContextVar("bronze_context")


@contextmanager
def landing_context(ctx: BronzeContext) -> Iterator[BronzeContext]:
    token = _context.set(ctx)
    try:
        yield ctx
    finally:
        _context.reset(token)


def land_bronze(run_id: str) -> str:
    """Stage: run_id -> sha256 -> staged file -> bronze rows, landed once per file."""
    ctx = _context.get()
    store = SnowflakeRunStore(ctx.session) if ctx.store is None else ctx.store
    run = store.get_run(run_id)
    if run is None:
        raise KeyError(run_id)
    stage_file = f"{run.sha256}.xlsx"
    data = ctx.read_file(stage_file)
    landed = LandedFile(
        file_sha256=run.sha256,
        filename=run.filename,
        stage_file=stage_file,
        size_bytes=len(data),
        first_run_id=run_id,
    )
    if ctx.land is not None:
        return ctx.land(landed, data)
    return land_file(ctx.session, landed, data, batch_size=ctx.batch_size)


def register_stages() -> None:
    """Idempotent: adds land_bronze once (D4)."""
    if STAGE_NAME not in registry.stages:
        registry.stages.add(STAGE_NAME, land_bronze)
